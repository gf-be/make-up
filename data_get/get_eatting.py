import argparse
import io
import json
import logging
import re
import subprocess
import sys
import time
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple
from urllib.parse import urljoin, urlparse


import openpyxl
import requests
from bs4 import BeautifulSoup, Tag

from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

try:
    from DrissionPage import ChromiumOptions, ChromiumPage
except ImportError:
    ChromiumOptions = None
    ChromiumPage = None


PROJECT_ROOT = Path(__file__).resolve().parent.parent
NODE_ATTACHMENT_PARSER = Path(__file__).resolve().parent / 'parse_announcement_attachment.js'
BASE_URL = 'https://zwfw.samr.gov.cn'
LIST_URL = f'{BASE_URL}/scjg/wyk/tbtg/'
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent / 'output'
DEFAULT_ITEMS_DIR = DEFAULT_OUTPUT_DIR / 'items'
DEFAULT_DOWNLOAD_DIR = Path(__file__).resolve().parent / 'downloads' / 'samr_food'
DEFAULT_BACKEND_URL = 'http://127.0.0.1:3000'
USER_AGENT = (
    'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
    '(KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36'
)
SUPPORTED_ATTACHMENT_EXTENSIONS = {'.xlsx', '.xls', '.doc', '.docx', '.pdf', '.zip'}
EXCEL_EXTENSIONS = {'.xlsx', '.xls'}
ARCHIVE_EXTENSIONS = {'.zip'}

CONTENT_CONTAINER_SELECTORS = [
    '.TRS_Editor',
    '.trs_editor_view',
    '.wp_articlecontent',
    '.article-content',
    '.detail-content',
    '.Custom_UnionStyle',
    'article',
    'main',
    '.content',
]
CONTENT_REMOVE_SELECTORS = [
    'script',
    'style',
    'noscript',
    'iframe',
    'form',
    'header',
    'footer',
    'nav',
    '.breadcrumb',
    '.crumb',
    '.share',
    '.shares',
    '.pages',
    '.pagination',
    '.related',
    '.attachment',
    '.attachments',
    '.editor-tools',
    '.article-source',
    '.article-info',
    '[class*="share"]',
    '[class*="breadcrumb"]',
    '[class*="footer"]',
    '[class*="header"]',
    '[class*="nav"]',
    '[class*="tool"]',
    '[class*="source"]',
]
CONTENT_NOISE_LINE_PATTERNS = [
    re.compile(r'^\|+$'),
    re.compile(r'^\d{4}年\d{1,2}月\d{1,2}日(?:\s*星期[一二三四五六日天])?$'),
    re.compile(r'^(发布时间|成文日期|文章来源|信息来源|来源|字号)[:：].{0,50}$'),
    re.compile(r'^(官方微信|官方微博|无障碍|长者版|进入关怀版|网站地图|联系我们|打印|关闭窗口|扫一扫在手机打开当前页)$'),
]
CONTENT_NAVIGATION_TOKENS = {
    '首页', '机构', '新闻', '政务', '服务', '互动', '数据', '专题',
    '知识产权', '质量强国', '总局', '要闻', '动态', '信息公开', '办事服务',
    '政务公开', '政策文件', '法规', '解读', '回应', '科普', '下载', '客户端',
}


HEADER_ALIASES = {
    '序号': 'sequence_no',
    '标称生产企业名称': 'company_names',
    '生产企业名称': 'company_names',
    '标称生产企业地址': 'company_addresses',
    '生产企业地址': 'company_addresses',
    '被抽样单位名称': 'sample_unit_name',
    '经营者名称': 'sample_unit_name',
    '销售单位名称': 'sample_unit_name',
    '网店名称': 'sample_unit_name',
    '被抽样单位地址': 'sample_unit_address',
    '经营地址': 'sample_unit_address',
    '样品名称': 'product_name',
    '产品名称': 'product_name',
    '规格型号': 'package_spec',
    '规格': 'package_spec',
    '商标': 'brand',
    '生产日期': 'production_date',
    '保质期': 'expiry_date',
    '不合格项目': 'unqualified_items',
    '检验值': 'inspection_value',
    '标准值': 'standard_value',
    '标签标注要求': 'label_requirement',
    '检验机构': 'inspection_institution',
    '食品细类': 'food_category',
    '抽样编号': 'sample_code',
    '备注': 'remarks',
    '省份': 'province',
}
REQUIRED_HEADER_KEYS = {'product_name', 'company_names', 'sample_unit_name', 'unqualified_items', 'inspection_institution'}


def build_session() -> requests.Session:
    session = requests.Session()
    retry = Retry(
        total=3,
        read=3,
        connect=3,
        backoff_factor=1,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset({'GET', 'HEAD', 'POST'}),
    )
    adapter = HTTPAdapter(max_retries=retry)
    session.mount('http://', adapter)
    session.mount('https://', adapter)
    session.headers.update({
        'User-Agent': USER_AGENT,
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
        'Cache-Control': 'no-cache',
        'Pragma': 'no-cache',
    })
    return session


def setup_logger(verbose: bool = False) -> logging.Logger:
    logger = logging.getLogger('samr_food_crawler')
    logger.setLevel(logging.DEBUG if verbose else logging.INFO)
    logger.handlers.clear()

    handler = logging.StreamHandler(sys.stdout)
    handler.setFormatter(logging.Formatter('[%(levelname)s] %(message)s'))
    handler.setLevel(logging.DEBUG if verbose else logging.INFO)
    logger.addHandler(handler)
    return logger


def clean_text(value: object) -> str:
    return str(value or '').replace('\u00a0', ' ').replace('\u3000', ' ').replace('\u0007', ' ').strip()


def normalize_whitespace(value: object) -> str:
    return re.sub(r'[ \t]+', ' ', clean_text(value).replace('\r', '\n')).strip()


def normalize_multiline_text(value: object) -> str:
    if value is None:
        return ''
    text = str(value).replace('\r\n', '\n').replace('\r', '\n')
    lines = [re.sub(r'[ \t]+', ' ', line).strip() for line in text.split('\n')]
    return '\n'.join(line for line in lines if line).strip()


def normalize_header(value: object) -> str:
    text = normalize_whitespace(value)
    return re.sub(r'[\s:：/（）()\[\]【】·]+', '', text)


def normalize_cell(value: object) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_whitespace(value)


def join_text(parts: Iterable[object], sep: str = '；', default: str = '') -> str:
    normalized = [normalize_whitespace(part) for part in parts if normalize_whitespace(part)]
    return sep.join(normalized) if normalized else default


def safe_filename(value: str, limit: int = 80) -> str:
    text = re.sub(r'[\\/:*?"<>|]+', '_', normalize_whitespace(value))
    text = re.sub(r'_+', '_', text).strip('._ ')
    return (text or 'announcement')[:limit]


def resolve_download_filename(attachment_name: str, attachment_url: str) -> str:
    raw_name = normalize_whitespace(attachment_name).replace('\\', '/').split('/')[-1]
    url_name = Path(urlparse(attachment_url).path).name
    source_name = raw_name or url_name or 'attachment.bin'
    source_path = Path(source_name)
    suffix = (source_path.suffix or Path(url_name).suffix or '.bin').lower()
    stem = source_path.stem if source_path.suffix else source_name
    return f'{safe_filename(stem, limit=100)}{suffix}'


def build_unique_path(path: Path) -> Path:
    if not path.exists():
        return path

    candidate = path
    counter = 1
    while candidate.exists():
        candidate = path.with_name(f'{path.stem}_{counter}{path.suffix}')
        counter += 1
    return candidate


def build_notice_storage_key(
    title: str,
    publish_date: Optional[str] = None,
    announcement_no: Optional[str] = None,
    detail_url: str = '',
) -> str:
    parts: List[str] = []
    if publish_date:
        parts.append(publish_date.replace('-', ''))
    if announcement_no:
        parts.append(safe_filename(announcement_no, limit=40))

    article_match = re.search(r'(art_[A-Za-z0-9]+)', detail_url or '')
    if article_match:
        parts.append(article_match.group(1))

    parts.append(safe_filename(title, limit=60))
    return '__'.join(part for part in parts if part)[:180]


def normalize_date_text(value: object) -> Optional[str]:

    text = normalize_whitespace(value)
    if not text:
        return None

    exact = re.search(r'(\d{4})-(\d{1,2})-(\d{1,2})', text)
    if exact:
        year, month, day = exact.groups()
        return f'{year}-{int(month):02d}-{int(day):02d}'

    chinese = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', text)
    if chinese:
        year, month, day = chinese.groups()
        return f'{year}-{int(month):02d}-{int(day):02d}'

    return None


def classify_food_notice(*values: object) -> Optional[str]:
    text = normalize_whitespace(' '.join(clean_text(value) for value in values if clean_text(value)))
    if not text or '食品' not in text or '抽检' not in text:
        return None
    if '不合格情况' in text and ('通报' in text or '通告' in text):
        return 'unqualified_sampling'
    if ('监督抽检' in text or '抽检情况' in text) and ('通报' in text or '通告' in text):
        return 'supervision_sampling'
    return None


def get_food_notice_label(notice_category: Optional[str]) -> str:
    return {
        'unqualified_sampling': '食品抽检不合格情况通报',
        'supervision_sampling': '食品监督抽检情况通报',
    }.get(notice_category or '', '待人工判断')


def is_food_announcement(title: str, extra_text: str = '') -> bool:

    return classify_food_notice(title, extra_text) is not None


def extract_announcement_no(title: str, content_text: str) -> Optional[str]:

    source_text = f'{title}\n{content_text[:1000]}'
    match = re.search(r'([\u4e00-\u9fa5A-Za-z]{0,20}〔\d{4}〕\d+号)', source_text)
    return match.group(1) if match else None


def extract_title(soup: BeautifulSoup) -> str:
    for selector in ('h1', '.arti-title', '.article-title', '.title'):
        element = soup.select_one(selector)
        if element:
            text = normalize_whitespace(element.get_text(' ', strip=True))
            if text:
                return text
    title_tag = soup.find('title')
    return normalize_whitespace(title_tag.get_text(' ', strip=True) if title_tag else '')


def extract_publish_date(soup: BeautifulSoup, fallback: Optional[str] = None) -> Optional[str]:
    if fallback:
        normalized = normalize_date_text(fallback)
        if normalized:
            return normalized

    page_text = soup.get_text('\n', strip=True)
    match = re.search(r'(发布时间|成文日期)[：:]?\s*(\d{4}-\d{1,2}-\d{1,2})', page_text)
    if match:
        return normalize_date_text(match.group(2))

    return normalize_date_text(page_text)


def is_content_noise_line(line: str) -> bool:
    normalized = normalize_whitespace(line)
    compact = normalized.replace(' ', '')
    if not compact:
        return True

    if any(pattern.match(compact) for pattern in CONTENT_NOISE_LINE_PATTERNS):
        return True

    tokens = [token for token in re.split(r'[\s|/、·•]+', normalized) if token]
    if tokens:
        joined = ''.join(tokens)
        if len(joined) <= 40 and all(token in CONTENT_NAVIGATION_TOKENS for token in tokens):
            return True

    if '官方微信' in compact or '官方微博' in compact:
        return True

    return False


def clean_content_lines(text: object) -> str:
    lines = normalize_multiline_text(text).split('\n')
    cleaned_lines: List[str] = []
    for line in lines:
        normalized = normalize_whitespace(line)
        if not normalized or is_content_noise_line(normalized):
            continue
        if cleaned_lines and cleaned_lines[-1] == normalized:
            continue
        cleaned_lines.append(normalized)
    return '\n'.join(cleaned_lines).strip()


def extract_content_candidate(node: Tag) -> str:
    fragment = BeautifulSoup(str(node), 'html.parser')
    for selector in CONTENT_REMOVE_SELECTORS:
        for child in fragment.select(selector):
            child.decompose()
    return clean_content_lines(fragment.get_text('\n', strip=True))


def extract_content_text(soup: BeautifulSoup) -> str:
    candidates: List[str] = []
    seen_texts = set()

    for selector in CONTENT_CONTAINER_SELECTORS:
        for node in soup.select(selector):
            text = extract_content_candidate(node)
            if text and text not in seen_texts:
                candidates.append(text)
                seen_texts.add(text)

    if candidates:
        return max(candidates, key=lambda text: (len(text), -len(text.split('\n'))))

    fallback_root = soup.body or soup
    return extract_content_candidate(fallback_root) or clean_content_lines(soup.get_text('\n', strip=True))



def select_list_items(soup: BeautifulSoup) -> List[Tag]:
    section_nodes = soup.select('div.list-section')
    for section in section_nodes:
        title_node = section.select_one('div.title a')
        if title_node and '通报通告' in normalize_whitespace(title_node.get_text(' ', strip=True)):
            section_items = list(section.select('ul.items li'))
            if section_items:
                return section_items

    fallback_selectors = [
        'div.seeList ul.gettonggao li',
        'ul.gettonggao li',
        'div.seeList li',
        'ul.items li',
    ]
    for selector in fallback_selectors:
        items = list(soup.select(selector))
        if items:
            return items
    return []


def extract_link_target(value: object) -> str:
    if isinstance(value, list):
        value = value[0] if value else ''
    text = clean_text(value)
    if not text or text.startswith('javascript:') or text.startswith('#'):
        return ''
    return text


def extract_next_page_url(soup: BeautifulSoup, current_url: str) -> Optional[str]:
    next_link = soup.select_one('a.pageNext')
    if next_link is None:
        for link in soup.select('a'):
            if '下一页' in normalize_whitespace(link.get_text(' ', strip=True)):
                next_link = link
                break

    if next_link is None:
        return None

    candidates = [
        next_link.get('tagname'),
        next_link.get('href'),
    ]
    onclick_text = clean_text(next_link.get('onclick'))
    onclick_match = re.search(r"queryArticleByCondition\(this,'([^']+)'\)", onclick_text)
    if onclick_match:
        candidates.append(onclick_match.group(1))

    for candidate in candidates:
        target = extract_link_target(candidate)
        if target:
            return urljoin(current_url, target)
    return None


def find_header_row(rows: List[List[str]]) -> Optional[int]:

    for index, row in enumerate(rows[:15]):
        header_keys = {HEADER_ALIASES.get(normalize_header(cell)) for cell in row}
        header_keys.discard(None)
        if len(header_keys & REQUIRED_HEADER_KEYS) >= 3 and 'product_name' in header_keys and 'unqualified_items' in header_keys:
            return index
    return None


def build_field_map(header_row: List[str]) -> Dict[str, int]:
    field_map: Dict[str, int] = {}
    for index, cell in enumerate(header_row):
        canonical = HEADER_ALIASES.get(normalize_header(cell))
        if canonical and canonical not in field_map:
            field_map[canonical] = index
    return field_map


def cell_by_field(row: List[str], field_map: Dict[str, int], field_name: str) -> str:
    index = field_map.get(field_name)
    if index is None or index >= len(row):
        return ''
    return normalize_cell(row[index])


def build_inspection_result(row: List[str], field_map: Dict[str, int]) -> str:
    inspection_value = cell_by_field(row, field_map, 'inspection_value')
    standard_value = cell_by_field(row, field_map, 'standard_value')
    label_requirement = cell_by_field(row, field_map, 'label_requirement')
    return join_text([
        f'检验值：{inspection_value}' if inspection_value else '',
        f'标准值：{standard_value}' if standard_value else '',
        f'标签标注要求：{label_requirement}' if label_requirement and label_requirement not in {'/', '-'} else '',
    ])


def build_requirement(row: List[str], field_map: Dict[str, int]) -> str:
    standard_value = cell_by_field(row, field_map, 'standard_value')
    label_requirement = cell_by_field(row, field_map, 'label_requirement')
    return join_text([
        f'标准值：{standard_value}' if standard_value else '',
        f'标签标注要求：{label_requirement}' if label_requirement and label_requirement not in {'/', '-'} else '',
    ], sep='\n')


def build_remarks(row: List[str], field_map: Dict[str, int]) -> str:
    return join_text([
        cell_by_field(row, field_map, 'remarks'),
        f'商标：{cell_by_field(row, field_map, "brand")}' if cell_by_field(row, field_map, 'brand') else '',
        f'食品细类：{cell_by_field(row, field_map, "food_category")}' if cell_by_field(row, field_map, 'food_category') else '',
        f'抽样编号：{cell_by_field(row, field_map, "sample_code")}' if cell_by_field(row, field_map, 'sample_code') else '',
        f'省份：{cell_by_field(row, field_map, "province")}' if cell_by_field(row, field_map, 'province') else '',
    ], sep='\n', default='/')


def append_distinct_text(existing: object, addition: object, sep: str = '；') -> str:
    values = []
    seen = set()
    for part in (existing, addition):
        text = normalize_multiline_text(part)
        if not text or text == '/':
            continue
        for token in [item.strip() for item in text.split(sep) if item.strip()]:
            if token not in seen:
                seen.add(token)
                values.append(token)
    return sep.join(values)


def is_continuation_issue_row(row: List[str], field_map: Dict[str, int]) -> bool:
    primary_fields = [
        'sequence_no',
        'product_name',
        'company_names',
        'sample_unit_name',
        'sample_unit_address',
        'package_spec',
        'production_date',
        'expiry_date',
        'inspection_institution',
    ]
    if any(cell_by_field(row, field_map, field) for field in primary_fields):
        return False

    return bool(
        cell_by_field(row, field_map, 'unqualified_items')
        or cell_by_field(row, field_map, 'inspection_value')
        or cell_by_field(row, field_map, 'standard_value')
        or cell_by_field(row, field_map, 'label_requirement')
        or cell_by_field(row, field_map, 'remarks')
    )


def merge_continuation_row(product_row: Dict[str, object], row: List[str], field_map: Dict[str, int]) -> None:
    merged_items = append_distinct_text(product_row.get('unqualified_items'), cell_by_field(row, field_map, 'unqualified_items'))
    merged_result = append_distinct_text(product_row.get('inspection_result'), build_inspection_result(row, field_map), sep='\n')
    merged_requirement = append_distinct_text(product_row.get('requirement'), build_requirement(row, field_map), sep='\n')
    merged_remarks = append_distinct_text(product_row.get('remarks'), cell_by_field(row, field_map, 'remarks'), sep='\n')

    if merged_items:
        product_row['unqualified_items'] = merged_items
    if merged_result:
        product_row['inspection_result'] = merged_result
    if merged_requirement:
        product_row['requirement'] = merged_requirement
    if merged_remarks:
        product_row['remarks'] = merged_remarks


def build_product_row(row: List[str], field_map: Dict[str, int], sequence_no: int) -> Optional[Dict[str, object]]:
    product_name = cell_by_field(row, field_map, 'product_name')
    if not product_name:
        return None

    requirement = build_requirement(row, field_map)
    return {
        'sequence_no': sequence_no,
        'product_name': product_name,
        'company_names': cell_by_field(row, field_map, 'company_names') or None,
        'company_addresses': cell_by_field(row, field_map, 'company_addresses') or None,
        'sample_unit_name': cell_by_field(row, field_map, 'sample_unit_name') or None,
        'sample_unit_address': cell_by_field(row, field_map, 'sample_unit_address') or None,
        'package_spec': cell_by_field(row, field_map, 'package_spec') or None,
        'batch_no': cell_by_field(row, field_map, 'sample_code') or None,
        'production_date': cell_by_field(row, field_map, 'production_date') or None,
        'expiry_date': cell_by_field(row, field_map, 'expiry_date') or None,
        'product_region': cell_by_field(row, field_map, 'province') or None,
        'registration_no': None,
        'production_license_no': None,
        'inspection_institution': cell_by_field(row, field_map, 'inspection_institution') or None,
        'unqualified_items': cell_by_field(row, field_map, 'unqualified_items') or None,
        'inspection_result': build_inspection_result(row, field_map) or None,
        'requirement': requirement or None,
        'remarks': build_remarks(row, field_map),
        'is_counterfeit': 0,
    }


def parse_excel_attachment_with_openpyxl(file_bytes: bytes, logger: logging.Logger) -> Tuple[List[Dict[str, object]], str]:

    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    all_rows: List[Dict[str, object]] = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        raw_rows = [[normalize_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        header_index = find_header_row(raw_rows)
        if header_index is None:
            logger.debug('工作表 %s 未识别到食品抽检表头，已跳过', sheet_name)
            continue

        field_map = build_field_map(raw_rows[header_index])
        blank_streak = 0
        sequence_no = 1
        current_product_row: Optional[Dict[str, object]] = None
        for row in raw_rows[header_index + 1:]:
            if not any(normalize_whitespace(cell) for cell in row):
                blank_streak += 1
                if blank_streak >= 3:
                    break
                continue

            blank_streak = 0
            product_row = build_product_row(row, field_map, sequence_no)
            if product_row:
                all_rows.append(product_row)
                current_product_row = product_row
                sequence_no += 1
                continue

            if current_product_row and is_continuation_issue_row(row, field_map):
                merge_continuation_row(current_product_row, row, field_map)

    message = '' if all_rows else '未从食品抽检附件中识别到可导入的表格明细。'
    return all_rows, message



def parse_excel_attachment_with_node(local_path: Path) -> Tuple[List[Dict[str, object]], str]:
    if not NODE_ATTACHMENT_PARSER.is_file():
        raise RuntimeError(f'未找到 Node 附件解析脚本: {NODE_ATTACHMENT_PARSER}')

    completed = subprocess.run(
        ['node', str(NODE_ATTACHMENT_PARSER), str(local_path)],
        cwd=str(PROJECT_ROOT),
        capture_output=True,
        text=True,
        encoding='utf-8',
        errors='replace',
        check=True,
    )
    output = completed.stdout.strip()
    if not output:
        raise RuntimeError('Node 附件解析器没有返回结果')

    data = json.loads(output)
    rows = data.get('rows') if isinstance(data, dict) else []
    if not isinstance(rows, list):
        rows = []
    normalized_rows = [
        row for row in rows
        if isinstance(row, dict) and normalize_whitespace(row.get('product_name'))
    ]
    message = clean_text(data.get('message') if isinstance(data, dict) else '')
    if not normalized_rows and not message:
        message = '未从 Excel 附件中识别到可导入的表格明细。'
    return normalized_rows, message


def parse_excel_attachment(
    file_bytes: bytes,
    logger: logging.Logger,
    local_path: Optional[Path] = None,
    file_ext: str = '.xlsx',
) -> Tuple[List[Dict[str, object]], str]:
    normalized_ext = (file_ext or '').lower()
    errors: List[str] = []

    if normalized_ext != '.xls':
        try:
            return parse_excel_attachment_with_openpyxl(file_bytes, logger)
        except Exception as exc:
            errors.append(f'openpyxl 解析失败: {exc}')
            logger.warning('openpyxl 解析 Excel 失败，准备回退 Node 解析器: %s', exc)

    if local_path is None:
        raise RuntimeError('缺少附件文件路径，无法回退 Node Excel 解析器')

    try:
        return parse_excel_attachment_with_node(local_path)
    except subprocess.CalledProcessError as exc:
        stderr = clean_text(exc.stderr)
        errors.append(f'Node 解析失败: {stderr or exc}')
    except Exception as exc:
        errors.append(f'Node 解析失败: {exc}')

    raise RuntimeError('；'.join(errors) or 'Excel 附件解析失败')


def parse_zip_attachment(local_path: Path, logger: logging.Logger) -> Dict[str, object]:
    extract_dir = build_unique_path(local_path.with_suffix(''))
    extract_dir.mkdir(parents=True, exist_ok=True)

    rows: List[Dict[str, object]] = []
    extracted_files: List[Dict[str, object]] = []
    messages: List[str] = []

    try:
        with zipfile.ZipFile(local_path) as archive:
            for member in archive.infolist():
                if member.is_dir():
                    continue

                member_name = Path(member.filename.replace('\\', '/')).name
                if not member_name:
                    continue

                member_ext = Path(member_name).suffix.lower()
                if member_ext not in EXCEL_EXTENSIONS:
                    continue

                extracted_path = build_unique_path(extract_dir / resolve_download_filename(member_name, member.filename))
                extracted_bytes = archive.read(member)
                extracted_path.write_bytes(extracted_bytes)

                try:
                    parsed_rows, message = parse_excel_attachment(
                        extracted_bytes,
                        logger,
                        local_path=extracted_path,
                        file_ext=member_ext,
                    )
                except Exception as exc:
                    parsed_rows = []
                    message = f'解析压缩包内 Excel 失败: {exc}'

                rows.extend(parsed_rows)
                extracted_files.append({
                    'file_name': member_name,
                    'local_path': str(extracted_path),
                    'file_ext': member_ext,
                    'parsedCount': len(parsed_rows),
                    'message': message,
                })
                if message:
                    messages.append(f'{member_name}: {message}')
    except zipfile.BadZipFile as exc:
        raise RuntimeError(f'ZIP 附件损坏或格式不受支持: {exc}') from exc

    if not extracted_files:
        message = '压缩包内未找到可解析的 Excel 明细附件。'
    elif rows:
        message = ''
    else:
        message = '；'.join(messages) if messages else '压缩包内 Excel 未识别到可导入数据。'

    return {
        'supported': True,
        'attachment_type': 'zip',
        'parsedCount': len(rows),
        'counterfeitCount': 0,
        'rows': rows,
        'message': message,
        'extracted_files': extracted_files,
    }


class SamrFoodCrawler:

    def __init__(
        self,
        max_pages: int,
        max_items: Optional[int],
        output_dir: Path,
        download_dir: Path,
        request_delay: float,
        logger: logging.Logger,
        browser_path: Optional[str] = None,
        browser_user_data_dir: Optional[str] = None,
        headless: bool = False,
    ) -> None:
        self.max_pages = max_pages
        self.max_items = max_items
        self.output_dir = output_dir
        self.download_dir = download_dir
        self.request_delay = max(request_delay, 0)
        self.logger = logger
        self.browser_path = browser_path or None
        self.browser_user_data_dir = browser_user_data_dir or None
        self.headless = headless
        self.session = build_session()
        self._page = None
        self.items_dir = output_dir / 'items'
        self.items_dir.mkdir(parents=True, exist_ok=True)
        self.download_dir.mkdir(parents=True, exist_ok=True)


    def close(self) -> None:
        if self._page is not None:
            try:
                self._page.quit()
            except Exception as exc:
                self.logger.debug('关闭 DrissionPage 失败: %s', exc)
            finally:
                self._page = None
        self.session.close()

    def ensure_browser(self):
        if self._page is not None:
            return self._page
        if ChromiumOptions is None or ChromiumPage is None:
            raise RuntimeError('未安装 DrissionPage，无法通过浏览器访问目标网站')

        options = ChromiumOptions()
        options.auto_port()
        options.headless(self.headless)
        options.set_load_mode('normal')
        options.set_timeouts(base=10, page_load=60, script=20)
        options.set_user_agent(USER_AGENT)
        options.set_download_path(str(self.download_dir))
        options.set_argument('--window-size=1440,900')
        options.set_argument('--lang=zh-CN')
        options.set_argument('--disable-blink-features=AutomationControlled')
        options.set_argument('--disable-features=Translate,AutomationControlled')
        options.set_argument('--disable-infobars')
        options.set_argument('--no-default-browser-check')
        options.set_argument('--disable-popup-blocking')
        options.set_argument('--disable-dev-shm-usage')
        options.set_argument('--ignore-certificate-errors')
        if self.browser_path:
            options.set_browser_path(self.browser_path)
        if self.browser_user_data_dir:
            user_data_dir = Path(self.browser_user_data_dir).resolve()
            user_data_dir.mkdir(parents=True, exist_ok=True)
            options.set_user_data_path(str(user_data_dir))

        self.logger.info(
            '启动 DrissionPage 浏览器: headless=%s, browser_path=%s, user_data_dir=%s',
            self.headless,
            self.browser_path or '<auto>',
            self.browser_user_data_dir or '<auto>',
        )
        self._page = ChromiumPage(options)
        return self._page

    def sync_browser_cookies_to_session(self) -> None:
        if self._page is None:
            return
        try:
            for cookie in self._page.cookies(all_info=True):
                name = cookie.get('name')
                value = cookie.get('value')
                if not name:
                    continue
                self.session.cookies.set(
                    name,
                    value,
                    domain=cookie.get('domain'),
                    path=cookie.get('path', '/'),
                )
        except Exception as exc:
            self.logger.debug('同步浏览器 Cookie 失败: %s', exc)

    def configure_browser_headers(self, referer: Optional[str] = None) -> None:
        page = self.ensure_browser()
        try:
            page.run_cdp('Network.enable')
            page.run_cdp(
                'Network.setExtraHTTPHeaders',
                headers={
                    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8',
                    'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
                    'Cache-Control': 'no-cache',
                    'Pragma': 'no-cache',
                    'Referer': referer or LIST_URL,
                    'Upgrade-Insecure-Requests': '1',
                },
            )
        except Exception as exc:
            self.logger.debug('配置浏览器请求头失败: %s', exc)

    def fetch_text(self, url: str, referer: Optional[str] = None) -> str:
        page = self.ensure_browser()
        self.configure_browser_headers(referer)
        page.get(url, retry=2, interval=2, timeout=60, show_errmsg=True)
        time.sleep(1.0)
        html = page.html or ''
        self.sync_browser_cookies_to_session()
        if not html.strip():
            raise RuntimeError(f'页面内容为空: {url}')
        return html

    def fetch_binary(self, url: str, referer: Optional[str] = None) -> bytes:
        self.sync_browser_cookies_to_session()
        headers = {
            'Referer': referer or LIST_URL,
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        }
        response = self.session.get(url, timeout=60, headers=headers)
        response.raise_for_status()
        return response.content


    def crawl(self) -> Dict[str, object]:
        self.logger.info('使用 DrissionPage 访问国家市场监督管理总局网站')
        candidates = self.load_candidates()
        results: List[Dict[str, object]] = []
        self.logger.info('识别到待抓取通报通告 %s 条', len(candidates))



        for index, item in enumerate(candidates, start=1):
            self.logger.info('[%s/%s] 抓取 %s', index, len(candidates), item['title'])
            try:
                result = self.process_item(item, index)
                results.append(result)
            except Exception as exc:
                self.logger.exception('处理通报通告失败: %s', item['detail_url'])

                results.append({
                    'sequence': index,
                    'title': item['title'],
                    'detail_url': item['detail_url'],
                    'publish_date': item.get('publish_date'),
                    'notice_category': item.get('notice_category'),
                    'classification_status': item.get('classification_status') or ('identified' if item.get('notice_category') else 'manual_review'),
                    'success': False,
                    'error': str(exc),
                })


            if self.request_delay > 0 and index < len(candidates):
                time.sleep(self.request_delay)

        manifest: Dict[str, object] = {
            'generated_at': datetime.now().isoformat(timespec='seconds'),
            'source': '国家市场监督管理总局-通报通告列表（全量抓取）',
            'list_url': LIST_URL,
            'items_dir': str(self.items_dir),
            'download_dir': str(self.download_dir),
            'total_candidates': len(candidates),
            'success_count': sum(1 for item in results if item.get('success')),
            'failed_count': sum(1 for item in results if not item.get('success')),
            'notice_category_counts': {
                'unqualified_sampling': sum(1 for item in results if item.get('notice_category') == 'unqualified_sampling'),
                'supervision_sampling': sum(1 for item in results if item.get('notice_category') == 'supervision_sampling'),
                'manual_review': sum(1 for item in results if item.get('classification_status') == 'manual_review'),
            },
            'results': results,
        }


        manifest_path = self.output_dir / 'samr_food_manifest.json'
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        self.logger.info('抓取完成，摘要已写入: %s', manifest_path)
        return manifest

    def load_candidates(self) -> List[Dict[str, object]]:
        candidates: List[Dict[str, object]] = []

        seen_urls = set()
        visited_pages = set()
        page = 1
        page_url = LIST_URL

        while page <= self.max_pages and page_url and page_url not in visited_pages:
            visited_pages.add(page_url)
            html = self.fetch_text(page_url, referer=LIST_URL)
            soup = BeautifulSoup(html, 'html.parser')
            page_items: List[Dict[str, object]] = []
            for li in select_list_items(soup):
                link = li.select_one('h3 a') or li.select_one('a[href]')
                if not link:
                    continue
                title = normalize_whitespace(link.get('title') or link.get_text(' ', strip=True))
                href = urljoin(page_url, extract_link_target(link.get('href')))
                notice_category = classify_food_notice(title, li.get_text(' ', strip=True))
                if not title or not href or href in seen_urls:
                    continue

                time_node = li.select_one('.time')
                publish_date = normalize_date_text(time_node.get_text(' ', strip=True) if time_node else li.get_text(' ', strip=True))
                candidate: Dict[str, object] = {
                    'title': title,
                    'detail_url': href,
                    'publish_date': publish_date or '',
                    'source_page': page_url,
                    'notice_category': notice_category,
                    'notice_category_label': get_food_notice_label(notice_category),
                    'classification_status': 'identified' if notice_category else 'manual_review',
                }

                candidates.append(candidate)
                page_items.append(candidate)
                seen_urls.add(href)

                if self.max_items and len(candidates) >= self.max_items:
                    return candidates

            next_page_url = extract_next_page_url(soup, page_url)
            self.logger.info(
                '列表第 %s 页新增候选 %s 条（待人工判断 %s 条），下一页：%s',
                page,
                len(page_items),
                sum(1 for item in page_items if item.get('classification_status') == 'manual_review'),
                next_page_url or '<无>'
            )
            if not next_page_url:
                break
            page += 1
            page_url = next_page_url

        return candidates


    def process_item(self, item: Dict[str, object], sequence: int) -> Dict[str, object]:

        detail_url = normalize_whitespace(item.get('detail_url'))
        source_page = normalize_whitespace(item.get('source_page')) or LIST_URL
        item_title = normalize_whitespace(item.get('title'))
        publish_date_hint = normalize_whitespace(item.get('publish_date')) or None

        html = self.fetch_text(detail_url, referer=source_page)
        soup = BeautifulSoup(html, 'html.parser')
        title = extract_title(soup) or item_title
        content_text = extract_content_text(soup)
        publish_date = extract_publish_date(soup, publish_date_hint)
        announcement_no = extract_announcement_no(title, content_text)
        notice_category = classify_food_notice(title, content_text, item.get('notice_category'))
        classification_status = 'identified' if notice_category else 'manual_review'
        notice_category_label = get_food_notice_label(notice_category)
        storage_key = build_notice_storage_key(title, publish_date, announcement_no, detail_url)
        attachments = self.extract_attachments(soup, detail_url, storage_key)


        payload = {
            'sequence': sequence,
            'title': title,
            'announcement_no': announcement_no,
            'publish_date': publish_date,
            'detail_url': detail_url,
            'source_page': source_page,

            'product_type': 'food',
            'announcement_type': 'sampling',
            'notice_category': notice_category,
            'notice_category_label': notice_category_label,
            'classification_status': classification_status,
            'requires_manual_review': classification_status == 'manual_review',
            '_import_meta': {
                'product_type': 'food',
                'announcement_type': 'sampling',
                'product_type_label': '食品',
                'announcement_type_label': '抽检通告',
                'notice_category': notice_category,
                'notice_category_label': notice_category_label,
                'classification_status': classification_status,
                'requires_manual_review': classification_status == 'manual_review',
            },
            'crawl_record': {
                'product_type': 'food',
                'announcement_type': 'sampling',
                'notice_category': notice_category,
                'classification_status': classification_status,
                'requires_manual_review': classification_status == 'manual_review',
                'source_page': source_page,
                'detail_url': detail_url,

            },
            'content_text': content_text,
            'content_preview': content_text[:1000],
            'attachments': attachments,
            'success': True,
        }

        publish_key = (publish_date or f'no-date-{sequence:03d}').replace('-', '')
        file_name = f'food_{publish_key}_{safe_filename(title)}.json'
        output_path = self.items_dir / file_name
        output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
        self.logger.info('已写入 JSON: %s', output_path)

        parsed_count = 0
        for attachment in attachments:
            parse_result = attachment.get('parse_result') if isinstance(attachment, dict) else None
            if isinstance(parse_result, dict):
                parsed_count += int(parse_result.get('parsedCount') or 0)

        return {
            'sequence': sequence,
            'title': title,
            'detail_url': detail_url,
            'publish_date': publish_date,
            'announcement_no': announcement_no,

            'notice_category': notice_category,
            'classification_status': classification_status,
            'attachment_count': len(attachments),
            'parsed_count': parsed_count,
            'output_file': str(output_path),
            'success': True,
        }



    def extract_attachments(self, soup: BeautifulSoup, detail_url: str, storage_key: str) -> List[Dict[str, object]]:
        attachments = []
        seen_urls = set()
        for link in soup.select('a[href]'):
            href = urljoin(detail_url, extract_link_target(link.get('href')))

            ext = Path(urlparse(href).path).suffix.lower()
            if not href or href in seen_urls or ext not in SUPPORTED_ATTACHMENT_EXTENSIONS:
                continue

            attachment_name = normalize_whitespace(link.get_text(' ', strip=True)) or Path(urlparse(href).path).name
            attachment_record = self.download_and_parse_attachment(storage_key, attachment_name, href, detail_url)
            attachments.append(attachment_record)
            seen_urls.add(href)
        return attachments

    def download_and_parse_attachment(
        self,
        storage_key: str,
        attachment_name: str,
        attachment_url: str,
        referer_url: str,
    ) -> Dict[str, object]:
        folder = self.download_dir / storage_key
        folder.mkdir(parents=True, exist_ok=True)
        local_path = build_unique_path(folder / resolve_download_filename(attachment_name, attachment_url))

        file_bytes = self.fetch_binary(attachment_url, referer=referer_url)
        local_path.write_bytes(file_bytes)
        file_ext = local_path.suffix.lower()

        parse_result: Dict[str, object]
        if file_ext in EXCEL_EXTENSIONS:
            try:
                rows, message = parse_excel_attachment(file_bytes, self.logger, local_path=local_path, file_ext=file_ext)
                parse_result = {
                    'supported': True,
                    'attachment_type': 'excel',
                    'parsedCount': len(rows),
                    'counterfeitCount': 0,
                    'rows': rows,
                    'message': message,
                }
            except Exception as exc:
                parse_result = {
                    'supported': True,
                    'attachment_type': 'excel',
                    'parsedCount': 0,
                    'counterfeitCount': 0,
                    'rows': [],
                    'message': f'解析食品抽检附件失败: {exc}',
                }
        elif file_ext in ARCHIVE_EXTENSIONS:
            try:
                parse_result = parse_zip_attachment(local_path, self.logger)
            except Exception as exc:
                parse_result = {
                    'supported': True,
                    'attachment_type': 'zip',
                    'parsedCount': 0,
                    'counterfeitCount': 0,
                    'rows': [],
                    'message': f'解析食品抽检压缩包失败: {exc}',
                }
        else:
            parse_result = {
                'supported': False,
                'attachment_type': file_ext.replace('.', '') or 'unknown',
                'parsedCount': 0,
                'counterfeitCount': 0,
                'rows': [],
                'message': '当前附件保留下载，不参与食品表格解析。',
            }

        raw_parsed_count = parse_result.get('parsedCount')
        parsed_count = int(raw_parsed_count) if isinstance(raw_parsed_count, (int, float, str)) and str(raw_parsed_count).strip() else 0
        parse_message = clean_text(parse_result.get('message'))


        return {
            'attachment_name': attachment_name,
            'attachment_url': attachment_url,
            'local_path': str(local_path),
            'file_ext': file_ext,
            'download': {
                'success': True,
                'bytes': len(file_bytes),
            },
            'parse_result': parse_result,
            'parse_error': None if parsed_count > 0 else (parse_message or None),
        }



def import_to_backend(
    backend_url: str,
    confirm_published: bool,
    logger: logging.Logger,
) -> Dict[str, object]:
    session = build_session()
    base_url = backend_url.rstrip('/')
    import_url = f'{base_url}/api/announcement-staging/import-json'
    confirm_url = f'{base_url}/api/announcement-staging/confirm-all'

    logger.info('开始将食品抽检情况 JSON 导入后端临时表: %s', import_url)

    import_response = session.post(
        import_url,
        json={'product_type': 'food', 'announcement_type': 'sampling'},
        timeout=180,
    )
    import_response.raise_for_status()
    import_result = import_response.json()
    if import_result.get('success') is False:
        raise RuntimeError(import_result.get('message') or '导入后端临时表失败')

    created_ids = [
        int(item['id'])
        for item in ((import_result.get('data') or {}).get('items') or [])
        if item.get('action') == 'created' and item.get('id')
    ]

    confirmed = None
    if confirm_published and created_ids:
        logger.warning('开始将新增食品抽检批次直接导入正式库（通常建议先在前端人工确认）: %s', created_ids)
        confirm_response = session.post(confirm_url, json={'ids': created_ids}, timeout=180)

        confirm_response.raise_for_status()
        confirmed = confirm_response.json()
        if confirmed.get('success') is False:
            raise RuntimeError(confirmed.get('message') or '导入正式库失败')
    elif confirm_published:
        logger.info('本次没有新增临时批次，跳过正式库导入')

    return {
        'import': import_result,
        'confirm': confirmed,
        'created_ids': created_ids,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description='使用 DrissionPage 抓取市场监管总局通报通告列表，尽可能解析食品相关附件，并输出为系统可导入 JSON')


    parser.add_argument('--max-pages', type=int, default=3, help='最多抓取列表页数，默认 3')
    parser.add_argument('--max-items', type=int, default=0, help='最多抓取通告数量，0 表示不限制')
    parser.add_argument('--output-dir', type=str, default=str(DEFAULT_OUTPUT_DIR), help='输出目录，默认 data_get/output')
    parser.add_argument('--download-dir', type=str, default=str(DEFAULT_DOWNLOAD_DIR), help='附件下载目录')
    parser.add_argument('--request-delay', type=float, default=0.8, help='请求之间的间隔秒数，默认 0.8')
    parser.add_argument('--backend-url', type=str, default=DEFAULT_BACKEND_URL, help='后端服务地址，默认 http://127.0.0.1:3000')
    parser.add_argument('--import-staging', action='store_true', help='抓取后自动导入后端临时表，默认仍由前端人工判断是否入正式库')
    parser.add_argument('--confirm-published', action='store_true', help='抓取后直接导入正式抽检通告模块（隐含 --import-staging，不建议日常使用）')

    parser.add_argument('--headless', action='store_true', help='以无头模式运行 DrissionPage，默认关闭以降低反爬风险')
    parser.add_argument('--browser-path', type=str, default='', help='Chromium/Chrome 可执行文件路径，可选')
    parser.add_argument('--user-data-dir', type=str, default='', help='浏览器用户数据目录，可选，用于复用缓存或登录态')
    parser.add_argument('--verbose', action='store_true', help='输出更详细日志')
    return parser.parse_args()



def main() -> int:
    args = parse_args()
    logger = setup_logger(args.verbose)
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    items_dir = DEFAULT_ITEMS_DIR if output_dir == DEFAULT_OUTPUT_DIR else output_dir / 'items'
    items_dir.mkdir(parents=True, exist_ok=True)
    crawler = SamrFoodCrawler(
        max_pages=max(args.max_pages, 1),
        max_items=args.max_items or None,
        output_dir=output_dir,
        download_dir=Path(args.download_dir).resolve(),
        request_delay=args.request_delay,
        logger=logger,
        browser_path=args.browser_path,
        browser_user_data_dir=args.user_data_dir,
        headless=args.headless,
    )


    try:
        manifest = crawler.crawl()
        notice_category_counts = manifest.get('notice_category_counts')
        raw_manual_review_count = notice_category_counts.get('manual_review', 0) if isinstance(notice_category_counts, dict) else 0
        manual_review_count = int(raw_manual_review_count) if isinstance(raw_manual_review_count, (int, float, str)) and str(raw_manual_review_count).strip() else 0

        logger.info(
            '通报通告抓取完成：成功 %s 条，失败 %s 条，待人工判断 %s 条，JSON 输出目录 %s',
            manifest['success_count'],
            manifest['failed_count'],
            manual_review_count,
            items_dir,
        )




        if args.import_staging or args.confirm_published:
            backend_result = import_to_backend(args.backend_url, args.confirm_published, logger)
            created_ids = backend_result.get('created_ids')
            if not isinstance(created_ids, list):
                created_ids = []
            logger.info('后端导入完成，新增临时批次 %s 个', len(created_ids))
            if backend_result.get('confirm'):
                logger.info('已同步导入正式抽检通告模块')


        return 0
    except Exception as exc:
        logger.exception('食品抽检通报抓取失败: %s', exc)
        return 1
    finally:
        crawler.close()


if __name__ == '__main__':
    raise SystemExit(main())
